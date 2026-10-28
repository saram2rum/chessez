#!/usr/bin/env python3
"""Oracle 섹션 3장 — PART 1 (Oracle 잠재력) 슬라이드 복붙용.

Slide 1: Oracle 정책 — 정의 및 best_t 측정
Slide 2: best_t 분포 — 평균 vs 개별 (통합 표)
Slide 3: Oracle 정책의 잠재력 — 38% 개선
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT = Path(__file__).resolve().parent.parent / "oracle_section.pptx"


def add_title(slide, text, size=30):
    title = slide.shapes.title
    title.text = text
    for p in title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True


def add_table(slide, left_in, top_in, width_in, height_in, headers, rows,
              header_font_size=14, body_font_size=14,
              highlight_cells=None, col_widths_in=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(left_in), Inches(top_in),
        Inches(width_in), Inches(height_in),
    )
    table = table_shape.table
    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            if i < n_cols:
                table.columns[i].width = Inches(w)
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(header_font_size)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(body_font_size)
    if highlight_cells:
        for r_idx, c_idx, color_hex in highlight_cells:
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(color_hex)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.bold = True
    return table


def add_textbox(slide, left_in, top_in, width_in, height_in, text,
                font_size=16, bold=False, italic=False, color_hex=None):
    tb = slide.shapes.add_textbox(
        Inches(left_in), Inches(top_in),
        Inches(width_in), Inches(height_in),
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.italic = italic
            if color_hex:
                r.font.color.rgb = RGBColor.from_string(color_hex)
    return tb


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # ─────────────────────────────────────────────────────────────
    # Slide 1: Oracle 정책 정의 (단순)
    # ─────────────────────────────────────────────────────────────
    slide1 = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide1, "Oracle 정책 — Proposed 의 이론적 상한")

    add_textbox(
        slide1, 0.5, 1.1, 12.3, 0.5,
        "'정답을 알면 어디까지 갈 수 있나?' — Proposed 차등 할당의 절대 상한 정량화",
        font_size=14, italic=True,
    )

    # Oracle 정책 정의
    add_textbox(
        slide1, 0.7, 1.9, 12.0, 0.5,
        "[Oracle 정책 정의]",
        font_size=18, bold=True,
    )
    add_textbox(
        slide1, 0.7, 2.4, 12.0, 2.0,
        "• 각 FEN i 의 best_t = argmin_t |예측(t) − GT|\n"
        "     → 그 FEN 의 오차를 최소화하는 시간 (50~500ms 중)\n\n"
        "• Oracle 정책 = 각 FEN 에 자기 best_t 부여",
        font_size=14,
    )

    # 특징
    add_textbox(
        slide1, 0.7, 4.6, 12.0, 0.5,
        "[특징]",
        font_size=18, bold=True,
    )
    add_textbox(
        slide1, 0.7, 5.1, 12.0, 2.0,
        "• Ground truth 알아야 계산 가능 → 이론적 상한 (oracle)\n"
        "• 어떤 정책도 Oracle 의 MAE 보다 낮을 수 없음\n"
        "• Proposed 차등 할당의 절대 상한 (= 잠재력의 최댓값)",
        font_size=14,
    )

    add_textbox(
        slide1, 0.5, 7.0, 12.3, 0.3,
        "→ 다음: best_t 의 분포 + Oracle 의 정량적 잠재력",
        font_size=12, italic=True,
    )

    # ─────────────────────────────────────────────────────────────
    # Slide 2: best_t 분포 — 평균 vs 개별 (카운트 방법 포함)
    # ─────────────────────────────────────────────────────────────
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide2, "best_t 분포 — 평균 vs 개별 위치")

    add_textbox(
        slide2, 0.5, 1.05, 12.3, 0.5,
        "전체 평균 (wu MAE) 과 개별 위치 (best_t) 는 다른 정보를 보여준다",
        font_size=14, italic=True,
    )

    # 카운트 방법 안내 (작은 글씨)
    add_textbox(
        slide2, 0.5, 1.5, 12.3, 0.4,
        "[best_t 분포 카운트 방법]  ① best_t 비율: 동률 시 가장 짧은 시간만 카운트 (합=100%)  "
        "/  ② 비중: 동률 시 모두 카운트 (합>100%)",
        font_size=11, italic=True,
    )

    # 통합 표
    SEARCHTIMES = [50, 100, 150, 200, 250, 300, 350, 400, 450, 500]
    wu_mae = [2.4201, 2.1595, 1.9941, 1.8776, 1.8029, 1.7349, 1.6791, 1.6250, 1.5865, 1.5396]
    best_t_ratio = [19.1, 12.1, 9.5, 9.6, 8.1, 7.9, 8.0, 8.8, 8.1, 8.9]
    bichung = [19.10, 18.35, 19.41, 21.97, 23.78, 25.90, 28.49, 33.45, 36.35, 41.89]

    rows = []
    for i, t in enumerate(SEARCHTIMES):
        rows.append([
            f"{t} ms",
            f"{wu_mae[i]:.4f}",
            f"{best_t_ratio[i]:.1f}%",
            f"{bichung[i]:.2f}%",
        ])

    # 강조: 50ms 의 best_t 비율, 500ms 의 wu MAE, 500ms 의 비중
    add_table(
        slide2,
        left_in=0.5, top_in=2.0, width_in=8.5, height_in=4.8,
        headers=["검색 시간", "wu MAE", "best_t 비율", "비중"],
        rows=rows,
        header_font_size=14, body_font_size=12,
        highlight_cells=[
            (1, 2, "B3D9FF"),    # 50ms best_t 비율
            (10, 1, "B3FFB3"),   # 500ms wu MAE
            (10, 3, "FFD9B3"),   # 500ms 비중
        ],
        col_widths_in=[2.0, 2.0, 2.3, 2.2],
    )

    # 우측 관찰
    add_textbox(
        slide2, 9.3, 2.0, 3.7, 0.5,
        "관찰",
        font_size=16, bold=True,
    )
    add_textbox(
        slide2, 9.3, 2.5, 3.7, 4.8,
        "▸ wu MAE 와 비중\n"
        "   둘 다 500ms 1위\n"
        "   시간 길수록 평균/분포\n"
        "   모두 우월 (자연스러움)\n\n"
        "▸ best_t 비율의 50ms 1위\n"
        "   (19.1%) 는 동률 처리에서\n"
        "   짧은 시간이 인덱스 우선\n"
        "   카운트되는 효과\n\n"
        "→ 시간↑ 우월은 평균/분포\n"
        "   모두 일관\n"
        "→ 그러나 다수 FEN 이 동률\n"
        "   (여러 시간에 비슷한 오차)\n"
        "→ 이 동률 분포가 차등 배분\n"
        "   잠재력의 근거",
        font_size=11,
    )

    # ─────────────────────────────────────────────────────────────
    # Slide 3: Oracle 정책의 잠재력
    # ─────────────────────────────────────────────────────────────
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide3, "Oracle 정책의 잠재력 — wu 대비 38% 개선")

    add_textbox(
        slide3, 0.5, 1.1, 12.3, 0.5,
        "Oracle 이 평균 243ms 만 쓰고도 시간 2배의 wu(500) 까지 압도 — 시간 vs 배분의 차이",
        font_size=14, italic=True,
    )

    # 비교 표
    rows_compare = [
        ["wu(50)",   "모든 FEN 50ms",   "50 ms",   "2.4201", "−53.8%",  "단순 위치만 OK, 나머지 큰 손해"],
        ["wu(250)",  "모든 FEN 250ms",  "250 ms",  "1.8029", "−37.9%",  "baseline (예산 중간)"],
        ["wu(500)",  "모든 FEN 500ms",  "500 ms",  "1.5396", "−27.4%",  "시간 2배여도 Oracle 못 따라잡음"],
        ["Oracle",   "각 FEN best_t",   "243 ms",  "1.1185", "기준",    "이론적 상한 (각 FEN 자기 best_t)"],
    ]

    add_table(
        slide3,
        left_in=0.5, top_in=1.8, width_in=12.3, height_in=3.0,
        headers=["정책", "정의", "평균 시간", "MAE", "vs Oracle", "비고"],
        rows=rows_compare,
        header_font_size=14, body_font_size=13,
        highlight_cells=[
            (1, 4, "FFB3B3"),       # wu(50) -53.8%
            (3, 4, "FFE0B3"),       # wu(500) -27.4%
            (4, 0, "B3FFB3"), (4, 1, "B3FFB3"), (4, 2, "B3FFB3"),
            (4, 3, "B3FFB3"), (4, 4, "B3FFB3"), (4, 5, "B3FFB3"),  # Oracle row green
        ],
        col_widths_in=[1.5, 2.2, 1.5, 1.5, 1.5, 4.1],
    )

    # 핵심 발견
    add_textbox(
        slide3, 0.5, 5.0, 12.3, 0.6,
        "[핵심 발견]",
        font_size=18, bold=True,
    )
    add_textbox(
        slide3, 0.5, 5.5, 12.3, 1.8,
        "▸ Oracle (243ms) vs wu(250): 같은 예산에서 38% 개선\n"
        "▸ Oracle (243ms) vs wu(500): 시간 절반 쓰고도 27% 더 정확\n"
        "  → 핵심: '시간 늘리기' 보다 '잘 배분하기' 가 효율적\n\n"
        "→ 차등 배분의 이론적 잠재력 명확 (Oracle 38% 개선)\n"
        "→ 다음 질문: 정적 FEN 정보만으로 이 잠재력에 도달 가능한가?",
        font_size=13,
    )

    prs.save(OUT)
    print(f"슬라이드 3장 생성: {OUT}")


if __name__ == "__main__":
    main()
