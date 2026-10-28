#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""학부 졸업논문 생성 — 윤지호 (202102296).

정보통신공학과 학부졸업논문 양식(.docx)을 베이스로 본문을 채워 최종 논문 생성.
내용은 최종발표.pptx를 정본으로 작성. (중간보고서의 S/M/L 버킷 실험은 제외,
재현성·데이터셋 등 현재에도 유효한 방법론만 참고하여 복원.)
현재 기준: 10개 시간 격자 {50..500}ms 차등 배분 vs 균일(wu),
이론적 상한(per-FEN best_t), 단변량(Spearman)·다변량(RandomForest) 예측 가능성.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE = "/Users/raon/Downloads/정보통신공학과_학부졸업논문양식 (1).docx"
OUT = "/Users/raon/Desktop/졸업논문/최종논문_윤지호.docx"


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith('}p') or child.tag.endswith('}tbl'):
            body.remove(child)


def run(p, text, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.font.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def body(doc, text, size=10, align=None, space_after=6, first_indent=True):
    p = doc.add_paragraph(style='Normal')
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.4
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(11)
    run(p, text, size=size)
    return p


def h1(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    run(p, text)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run(p, text)
    return p


def bullet(doc, text, size=10):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    run(p, "• " + text, size=size)
    return p


def caption(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run(p, text, bold=True, size=9)
    return p


def make_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ''
        pr = hdr[i].paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(pr, htxt, bold=True, size=9)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            pr = cells[i].paragraphs[0]
            pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(pr, str(val), size=9)
    return t


def title_page(doc):
    def cp(text, bold=False, size=11, before=0, after=6):
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        run(p, text, bold=bold, size=size)
        return p

    cp("학부 졸업논문", bold=True, size=13, before=24, after=30)
    cp("체스 형세 평가치 품질 향상을 위한", bold=True, size=18, after=2)
    cp("FEN 기반 엔진 탐색 시간 차등 배분의 검증과 한계", bold=True, size=18, after=2)
    cp(": t3.micro 실시간 서비스를 대상으로", bold=True, size=14, after=40)

    cp("대 학 : 공 과 대 학          학 과 : 정보통신공학과", size=12, after=8)
    cp("학 번 : 202102296          성 명 : 윤 지 호", size=12, after=44)

    cp("이 논문을 제출 하오니 승인하여 주십시오.", size=11, after=18)
    cp("2026 년    월    일", size=11, after=14)
    cp("성 명 : 윤 지 호 (인)", size=11, after=40)

    cp("위 학생의 논문 제출을 승인함.", size=11, after=18)
    cp("2026 년    월    일", size=11, after=14)
    cp("지 도 교 수 :              (인)", size=11, after=10)
    doc.add_page_break()


def main():
    doc = Document(TEMPLATE)
    clear_body(doc)

    title_page(doc)

    # ---------------- 요약 ----------------
    h1(doc, "요    약")
    body(doc,
         "본 논문은 저사양 클라우드 인스턴스(AWS EC2 t3.micro)에서 실시간 체스 형세 평가"
         "(Evaluation Bar) 서비스를 운영할 때, 한정된 연산 자원을 국면별로 차등 배분하여 평가치 "
         "품질을 향상시킬 수 있는지를 검증한다. 핵심 가설은 “국면 복잡도에 따라 엔진 탐색 시간을 "
         "차등 배분하면, 동일한 평균 시간 예산에서도 균일 배분보다 더 정확한 평가치를 얻을 수 있다”"
         "는 것이다. 이를 위해 각 국면(FEN)에 탐색 시간 격자 {50, 100, …, 500} ms 중 하나를 "
         "배정하는 차등 배분 정책과, 모든 국면에 동일한 시간을 주는 균일 정책(wu)을 동일 평균 예산 "
         "조건에서 비교하였다. 먼저 각 국면이 자신에게 최적인 탐색 시간(best_t)을 부여받는 이론적 "
         "상한(oracle) 정책을 측정한 결과, 이론적 상한은 평균 243ms만으로 동일 예산의 wu(250) 대비 "
         "약 38%, 시간을 두 배 쓰는 wu(500) 대비 약 27% 낮은 오차를 달성하여 차등 배분 전략이 "
         "원리적으로 큰 잠재력을 가짐을 확인하였다. 그러나 이 잠재력에 도달하려면 탐색 이전에 정적 "
         "정보(FEN)만으로 best_t를 예측할 수 있어야 한다. 7종의 정적 feature와 best_t 간 Spearman "
         "순위 상관은 모두 |ρ| < 0.05로 Cohen의 ‘small effect’ 임계값(0.10)의 절반에도 미치지 "
         "못하였고, 7종을 모두 입력한 RandomForest 회귀 역시 균일 정책과 통계적으로 구분되지 않는 "
         "성능(test MAE 1.811 대 1.810)에 그쳤다. 즉 차등 배분의 이론적 잠재력은 크지만, 정적 FEN "
         "정보만으로는 이를 실현할 수 없다. 본 연구는 negative result에 해당하나, 동일 예산 비교와 "
         "이론적 상한 실험을 통해 가설이 성립하지 않는 원인을 ‘차등 배분 전략 자체의 한계’가 아니라 "
         "‘정적 정보의 예측력 한계’로 분리하여 규명한 데 의의가 있다.",
         space_after=4)
    body(doc,
         "주제어(Keywords): 체스 엔진, Stockfish, 형세 평가치, 탐색 시간 차등 배분, 국면 복잡도, "
         "이론적 상한, RandomForest, negative result, t3.micro",
         space_after=10, first_indent=False)
    doc.add_page_break()

    # ---------------- 1. 서론 ----------------
    h1(doc, "1. 서론")
    body(doc,
         "체스 형세 평가치(evaluation)는 현재 국면이 어느 쪽에 얼마나 유리한지를 수치로 나타낸 "
         "값으로, 일반적으로 분석·복기 단계에서 활용된다. 본 연구의 배경이 되는 서비스 "
         "Chessez(https://chessez.com)는 사용자 간 1:1 실시간 대국 중 매 턴 형세 막대"
         "(Evaluation Bar)를 통해 평가치와 승률을 실시간으로 제공하는 변형 체스 서비스로, AWS "
         "EC2 t3.micro 단일 인스턴스 위에서 운영된다. 평가치는 백엔드 서버의 Stockfish 엔진이 "
         "산출하여 모든 사용자에게 동일하게 전달되므로, 서버 측에서 평가치를 계산하는 단일 진실 "
         "공급원(Single Source of Truth, SSoT) 구조를 갖는다.")
    body(doc,
         "이러한 실시간 서비스에서는 응답 지연이 사용자 경험을 직접 좌우한다. 본 연구는 "
         "Miller[1, p. 268]가 인간이 대기를 인식하기 시작하는 심리적 경계로 제시한 2초를 근거로, "
         "형세 분석 응답의 p95 꼬리 지연 상한을 2초로 설정한다. 한편 엔진의 탐색 시간이 길수록 "
         "평가치는 일반적으로 더 정확해진다. 따라서 제한된 서버 자원과 응답 지연 상한 아래에서 "
         "‘모든 국면에 동일한 고정 연산 시간을 부여하는’ 기존 방식은 비효율적일 수 있다. 어떤 "
         "국면은 짧은 탐색으로도 결론이 분명한 반면, 어떤 국면은 깊은 탐색이 필요하기 때문이다. "
         "예컨대 한 수 체크메이트 국면은 얕은 탐색으로도 충분하지만, 비대칭 폰 전개와 같이 깊은 "
         "탐색이 없으면 평가치가 뒤바뀔 수 있는 국면도 존재한다. 두 국면에 같은 연산 시간을 "
         "부여하는 것은 자원의 비효율적 사용이다.")
    body(doc,
         "이로부터 본 연구의 핵심 과제가 도출된다: 한정된 전체 연산 시간 예산을 국면별로 어떻게 "
         "분배해야, 실시간성을 유지하면서도 평가치의 품질을 높일 수 있는가? 구체적으로 본 연구는 "
         "“국면 복잡도에 따라 탐색 시간을 차등 배분하면, 동일한 평균 시간 예산에서도 균일 배분보다 "
         "더 정확한 평가치를 얻을 수 있다”는 가설을, 탐색 이전에 확보 가능한 정적 정보(FEN)만을 "
         "이용하는 관점에서 검증한다.")
    body(doc,
         "본 연구의 결론은 향상이 아닌 검증과 한계 규명에 가깝다. 본 논문의 기여는 다음과 같다. "
         "첫째, 각 국면이 자신의 최적 탐색 시간을 부여받는 이론적 상한 정책을 측정하여, 차등 배분 "
         "전략이 동일 예산 균일 정책 대비 약 38%의 개선 잠재력을 가짐을 정량적으로 보였다. 둘째, "
         "단변량 상관 분석(Spearman)과 다변량 기계학습(RandomForest)을 통해, 정적 FEN 정보만으로는 "
         "이 잠재력에 도달할 수 없음을 일관되게 확인하였다. 셋째, 동일 예산 비교와 이론적 상한 "
         "실험을 통해 negative result의 원인을 ‘전략의 한계’가 아닌 ‘정적 정보의 예측력 한계’로 "
         "분리함으로써, 후속 연구가 동일한 방향으로 자원을 투입하지 않도록 하는 가이드를 제공한다.")
    body(doc,
         "본 논문의 구성은 다음과 같다. 2장에서 관련 연구를 정리하고, 3장에서 서비스의 시스템 "
         "아키텍처를 기술한다. 4장에서 실험 설계와 방법(데이터셋, 평가 지표, 재현성 확보, 비교 정책, "
         "이론적 상한, 상관·기계학습 분석)을 설명하고, 5장에서 실험 결과를 제시·분석한다. 6장에서 "
         "결론과 향후 과제를 논의한다.")

    # ---------------- 2. 관련 연구 ----------------
    h1(doc, "2. 관련 연구")
    h2(doc, "2.1 수평 확장(Scale-Out)")
    body(doc,
         "가입자 2억 5천만 명 이상을 보유한 업계 1위 Chess.com은 수평 확장 방식을 채택한다[3]. "
         "2023년 트래픽 폭증 당시 약 200만 달러 규모의 하드웨어를 추가 투입하였으며[4], 현재 전 "
         "세계 여러 지역에 서버를 분산 배치하여 부하를 관리하고 지연을 최소화한다[5]. 이 방식은 "
         "수백만 동시 게임을 안정적으로 처리할 수 있으나, 막대한 예산이 전제되어야 한다. 본 연구는 "
         "스케일 아웃 없이 자원 제약적 단일 클라우드 인스턴스(t3.micro)를 전제로 애플리케이션 "
         "계층의 제어에 집중한다는 점에서 차별화된다.")
    h2(doc, "2.2 WASM 기반 브라우저 실행(Offloading)")
    body(doc,
         "업계 2위 Lichess는 Stockfish를 WebAssembly(WASM)로 포팅하여 분석 엔진을 클라이언트 "
         "브라우저에서 직접 실행한다[6]. 분석 연산이 서버가 아닌 사용자 기기에서 수행되므로 대규모 "
         "분석 요청을 서버 부담 없이 처리할 수 있다. 그러나 평가치가 클라이언트 환경에서 산출되기 "
         "때문에 단일 진실 공급원(SSoT)을 보장할 수 없다. 본 연구는 서버 사이드에서 평가치를 "
         "산출하는 SSoT 아키텍처를 유지하여 사용자 간 상태 불일치를 구조적으로 차단한다는 점에서 "
         "차별화된다.")
    h2(doc, "2.3 엔진 내부 타임 매니지먼트")
    body(doc,
         "Stockfish는 대국 모드(wtime/btime)에서 평가치 하락, 최선수 불안정성, 포지션 복잡도 등을 "
         "인자로 탐색 시간을 동적으로 조절하는 내부 타임 매니지먼트 메커니즘을 갖추고 있다. 그러나 "
         "이는 탐색을 시작한 뒤 내부 상태에 의존하여 계산된다. 반면 본 연구는 탐색 이전에 FEN만으로 "
         "복잡도를 사전 추정하여 탐색 시간을 결정한다는 점에서 근본적으로 다르다. 즉 본 연구가 묻는 "
         "것은 ‘탐색 전 정적 정보만으로 필요한 연산량을 예측할 수 있는가’이다.")
    h2(doc, "2.4 평가치–승률 변환과 평가 안정성")
    body(doc,
         "평가치의 오차를 측정할 때 센티폰(centipawn) 차이를 그대로 쓰는 것은 적절하지 않다. 동일한 "
         "센티폰 차이라도 승부가 팽팽한 0cp 근처에서는 승률에 큰 영향을 주지만, 한쪽이 크게 유리한 "
         "국면(±1000cp 이상)에서는 거의 영향이 없기 때문이다. 따라서 평가치 오차의 실질적 영향을 "
         "측정하려면 승률 기준 비교가 필요하며, 본 연구는 Lichess가 공개한 센티폰–승률 변환 공식을 "
         "사용한다[7]. 한편 엔진 평가치의 안정성 자체를 정량적으로 분석하려는 시도도 보고된 바 "
         "있다[9].")

    # ---------------- 3. 시스템 아키텍처 ----------------
    h1(doc, "3. 시스템 아키텍처")
    h2(doc, "3.1 서비스 개요 및 제약")
    body(doc,
         "대상 서비스 Chessez는 t3.micro 단일 인스턴스에서 실시간 대국과 형세 분석을 함께 제공한다. "
         "설계 제약(Constraint)은 형세 분석 응답의 p95 지연을 2초 이내로 유지하는 것이다[1]. "
         "문제(Issue)는 모든 국면에 고정된 연산 시간을 할당하는 기존 방식의 비효율성이며, "
         "목표(Objective)는 국면 복잡도에 따른 탐색 시간 차등 배분을 통해 동일 예산에서 평가치 "
         "품질을 개선하는 것이다.")
    h2(doc, "3.2 요청 처리 파이프라인")
    body(doc,
         "요청 처리는 세 단계로 구성된다. (1) FEN 분석·예산 할당: 요청을 수신하여 국면 복잡도를 "
         "분석하고 탐색 시간(movetime) 예산을 할당한다. (2) 대기열 방어: 동일 room_id 기준으로 기존 "
         "요청을 확인하고, 더 이상 유효하지 않은 오래된 요청을 폐기(Stale Request Drop)하여 대기열을 "
         "조정한다. (3) 엔진 연산·결과 반환: 할당된 예산으로 Stockfish 18이 탐색을 수행하되, 2초 "
         "예산 초과가 예상되면 해당 요청을 폐기하고, 그렇지 않으면 평가치를 반환한다. 이 구조는 "
         "실시간성(p95 ≤ 2초)을 보장하기 위한 방어 장치를 포함한다.")
    h2(doc, "3.3 복잡도 기반 예산 할당")
    body(doc,
         "차등 배분의 핵심은 탐색 이전에 FEN만으로 계산 가능한 정적 정보를 이용해 국면 복잡도를 "
         "추정하고, 복잡한 국면에 더 많은 탐색 시간을 부여하는 것이다. 탐색 시간 후보는 "
         "{50, 100, 150, …, 500} ms의 10개 격자로 정의하며, 평균 시간 예산이 동일한 균일 정책과 "
         "비교한다. 복잡도 추정에 사용 가능한 정적 feature로는 재료 불균형(imbalance), 합법 캡처 "
         "수(capture_count), 합법 수 총 개수(legal_moves), 비폰 재료 가치 합(total_material), "
         "양측 캡처·체크 수 합(forcing_total), 체크 여부(is_check), 종반 여부(is_endgame) 등이 "
         "있으며, 이들이 실제로 필요한 탐색 시간을 예측하는지가 본 연구의 검증 대상이다.")
    h2(doc, "3.4 기반 기술")
    body(doc,
         "요청 큐 관리와 복잡도 기반 예산 분배는 Python 기반 경량 비동기 프레임워크인 FastAPI로 "
         "구현하였다. 평가치 산출 엔진으로는 Stockfish 18을 사용하였다. Stockfish는 전 세계 체스 "
         "엔진 성능 1위로 Chess.com·Lichess 등 주요 플랫폼의 메인 엔진이며[11], 완전한 오픈소스로서 "
         "백엔드에서 명령어로 탐색 시간을 직접 제어하기에 적합하다.")

    # ---------------- 4. 실험 설계 및 방법 ----------------
    h1(doc, "4. 실험 설계 및 방법")
    h2(doc, "4.1 데이터셋")
    body(doc,
         "데이터셋은 일반 사용자 환경을 대표하도록 Lichess 래피드(rapid) 평균 레이팅 1200~1600 "
         "구간의 랜덤 50게임에서 추출하였다. 초기에는 단일 강자(Magnus Carlsen)의 게임을 사용했으나, "
         "한쪽이 일방적으로 유리한 국면에 분포가 편향될 수 있어 변경하였다. 중복 국면과 종료 국면을 "
         "제외한 총 2,927개의 FEN을 워크로드로 사용하였다.")
    h2(doc, "4.2 기준값(Ground Truth)")
    body(doc,
         "각 FEN의 정답 평가치는 Stockfish를 depth 20으로 탐색하여 산출하였다. Acher와 Esnault[2]는 "
         "depth 20을 ‘상당히 높은 깊이’로 설정하여 약 2억 7천만 개의 고유 포지션 평가 데이터베이스를 "
         "구축한 바 있으므로, depth 20은 충분히 높은 기준값으로 인정된다.")
    h2(doc, "4.3 평가 지표")
    body(doc,
         "평가 지표는 평균 절대 오차(Mean Absolute Error, MAE)를 사용하되, 엔진 평가치(centipawn)가 "
         "아니라 이를 승률로 변환한 값(0~100%)을 비교 대상으로 한다(2.4절)[7]. 즉 정책이 부여한 "
         "탐색 시간에서 산출한 평가치의 승률과 기준값(depth 20) 승률 간의 절대 오차를 전체 FEN에 "
         "대해 평균한다. MAE는 낮을수록 평가치 품질이 높음을 의미한다.")
    h2(doc, "4.4 재현성 확보")
    body(doc,
         "초기 실험에서는 동일 조건의 반복 실험에서도 평가치가 일관되게 산출되지 않아 MAE 자체가 "
         "흔들리는 문제가 관찰되었다. 원인으로는 Stockfish 병렬 탐색(Lazy SMP)[8]의 탐색 경로 "
         "비결정성, 공유 해시 테이블 잔상에 의한 연산 간섭, 클라우드 가상화 환경(t3.micro)의 미세한 "
         "성능 편차가 추정되었다. 이를 해소하기 위해 다음의 멱등성(동일 요청에 항상 동일 응답이 "
         "반환되는 성질) 확보 조치를 적용하였다.")
    bullet(doc, "탐색 환경 표준화: Threads=1로 직렬 탐색 고정, 매 요청 시 해시 테이블 초기화")
    bullet(doc, "하드웨어 비의존적 연산량 제어: movetime을 누적 노드 수(Nodes)로 치환. NPS(초당 탐색 노드 수)를 144회 측정하여 중앙값 약 364k를 얻었으며, 보수적으로 360,000 NPS(1ms = 360 nodes)를 변환 기준으로 채택")
    body(doc,
         "이 조치 이후 동일 요청에 대해 소수점까지 일치하는 멱등성을 확보하였고, 하드웨어에 독립적인 "
         "벤치마크 환경을 구축함으로써 MAE 측정의 정밀도를 확보하였다. 이후 모든 실험은 이 표준 "
         "환경에서 수행되었다.")
    h2(doc, "4.5 비교 정책")
    body(doc,
         "본 연구는 다음 두 정책을 동일한 평균 시간 예산 조건에서 비교한다. 탐색 시간 격자는 "
         "{50, 100, 150, …, 500} ms의 10개 값으로 정의한다.")
    bullet(doc, "Proposed(차등 할당): 각 FEN에 격자 중 하나의 탐색 시간을 배정한다. 평균 탐색 시간 T는 정책에 따라 결정되며, 본 연구는 이론적 상한·정적 feature 기반 예측 등 다양한 정책을 평가한다.")
    bullet(doc, "Baseline(균일 할당, wu): 모든 FEN에 동일한 시간 wu(T)를 적용한다. T는 Proposed의 평균 시간과 일치(또는 가장 가까운 격자)시켜 동일 예산 비교를 보장한다.")
    body(doc,
         "따라서 핵심 질문은 다음과 같다: “같은 평균 시간 예산(T ms) 안에서, 차등 할당이 균일 할당"
         "(wu)보다 더 정확한가?” 단순히 시간을 더 쓰면 MAE는 낮아지므로, 공정한 비교를 위해서는 "
         "반드시 동일 예산에서 비교해야 한다.")
    h2(doc, "4.6 이론적 상한(Oracle) 정의")
    body(doc,
         "차등 배분 전략이 도달 가능한 상한을 측정하기 위해, 각 FEN에 대해 격자 10개 시간 중 "
         "기준값과의 오차가 최소가 되는 시간 best_t를 부여하는 가상의 정책을 정의한다. best_t는 "
         "기준값(정답)을 알아야 계산할 수 있으므로 실제 서비스에서는 구현할 수 없는 이론적 상한이며, "
         "어떤 정책도 이 상한보다 낮은 MAE를 가질 수 없다. 이 상한과 균일 정책의 격차가 곧 차등 "
         "배분이 원리적으로 확보할 수 있는 최대 잠재력에 해당한다.")
    h2(doc, "4.7 정적 feature와 best_t의 상관 분석")
    body(doc,
         "이론적 상한의 잠재력에 정적 정보만으로 도달할 수 있는지 검토하기 위해, FEN에서 추출 가능한 "
         "7종의 정적 feature(3.3절)와 best_t 간의 Spearman 순위 상관 ρ를 분석한다. Spearman ρ는 두 "
         "값이 같은 방향으로 움직이는 정도를 −1~+1로 나타내며, 실제 값이 아니라 순위만 비교하므로 "
         "곡선 형태의 단조 관계도 포착하는 관대한 지표다. 상관의 크기 판단은 Cohen[10]의 기준"
         "(ρ ≈ 0.10 small, 0.30 medium, 0.50 large)을 따르며, |ρ| < 0.10은 예측 신호로서 유의미하지 "
         "않은(negligible) 것으로 본다.")
    h2(doc, "4.8 다변량 예측 가능성 검증(RandomForest)")
    body(doc,
         "단변량 상관이 약하더라도 여러 feature의 조합·비선형 관계로는 best_t를 예측할 수 있을지 "
         "모른다. 이를 검증하기 위해 7종 정적 feature를 모두 입력으로 RandomForest 회귀 모델을 "
         "학습시켜 best_t를 직접 예측하고, 예측한 시간을 격자에 스냅(snap)한 뒤 그 시간에서의 MAE를 "
         "측정하였다. RandomForest는 여러 결정트리의 예측을 평균내는 앙상블 모델로 비선형·feature 간 "
         "상호작용까지 자동으로 포착하므로, 신호가 조금이라도 존재한다면 이를 잡아낼 수 있는 강한 "
         "검증 수단이다(트리 400개, 잎 최소 5개). 검증은 전체 국면을 train 80% / test 20%로 무작위 "
         "분할하여 train으로 학습하고 보지 않은 test에서만 성능을 측정하였으며, 서로 다른 분할로 10회 "
         "반복하여 특정 분할 운에 따른 결과가 아님을 확인하였다.")

    # ---------------- 5. 실험 결과 및 분석 ----------------
    h1(doc, "5. 실험 결과 및 분석")
    h2(doc, "5.1 best_t 분포")
    body(doc,
         "각 FEN의 best_t 분포는 표 1과 같다. best_t는 동률(여러 시간이 동일한 최소 오차)일 때 가장 "
         "짧은 시간만 카운트하였다. 분포가 특정 시간에 집중되지 않고 10개 격자에 비교적 고르게 "
         "퍼져 있어, 국면마다 최적 탐색 시간이 서로 다름을 알 수 있다. 50ms의 비율이 상대적으로 "
         "높은 것은 동률 시 가장 짧은 시간을 카운트하는 규칙의 영향이다.")
    make_table(doc,
        ["탐색 시간", "best_t 비율", "탐색 시간", "best_t 비율"],
        [["50 ms", "19.1%", "300 ms", "7.9%"],
         ["100 ms", "12.1%", "350 ms", "8.0%"],
         ["150 ms", "9.5%", "400 ms", "8.8%"],
         ["200 ms", "9.6%", "450 ms", "8.1%"],
         ["250 ms", "8.1%", "500 ms", "8.9%"]])
    caption(doc, "표 1. best_t 분포 (동률 시 최단 시간 카운트, 합계 100%)")
    h2(doc, "5.2 이론적 상한의 잠재력")
    body(doc,
         "차등 배분 전략이 원리적으로 어느 정도의 효과를 낼 수 있는지를 먼저 확인하기 위해, 전체 "
         "2,927개 FEN에 대해 이론적 상한(각 FEN best_t)과 균일 정책의 MAE를 비교하였다(표 2). "
         "이론적 상한은 평균 243ms만을 사용하면서도 MAE 1.1185를 달성하여, 같은 예산의 "
         "wu(250)(MAE 1.8029)보다 약 38% 낮은 오차를 보였다. 더욱이 시간을 두 배 쓰는 "
         "wu(500)(MAE 1.5396)조차 이론적 상한을 따라잡지 못하여, 상한은 절반의 시간으로도 약 27% "
         "더 정확하였다.")
    make_table(doc,
        ["정책", "정의", "평균 시간", "MAE", "비고"],
        [["wu(250)", "모든 FEN 250ms", "250 ms", "1.8029", "동일 예산 baseline"],
         ["wu(500)", "모든 FEN 500ms", "500 ms", "1.5396", "시간 2배에도 상한 미달"],
         ["이론적 상한", "각 FEN best_t", "243 ms", "1.1185", "도달 불가능한 하한"]])
    caption(doc, "표 2. 이론적 상한과 균일 정책의 비교 (전체 2,927 FEN)")
    body(doc,
         "이 결과는 차등 배분 전략이 원리적으로 큰 잠재력(wu(250) 대비 약 38% 개선)을 가짐을 분명히 "
         "보여준다. 즉 차등 배분이라는 전략 자체에는 문제가 없다. 핵심 질문은 다음으로 이동한다: "
         "정답을 모르는 상태에서, 정적 FEN 정보만으로 이 잠재력에 도달할 수 있는가?")
    h2(doc, "5.3 단변량 상관 분석 결과")
    body(doc,
         "7종 정적 feature와 best_t 간 Spearman 순위 상관은 표 3과 같다. 모든 feature의 |ρ|가 0.05 "
         "미만으로, Cohen[10]의 ‘small effect’ 임계값 0.10의 절반에도 미치지 못하였다. 즉 어떤 단일 "
         "정적 feature도 best_t에 대한 유의미한 예측 신호를 제공하지 못한다(negligible).")
    make_table(doc,
        ["feature", "설명", "Spearman ρ"],
        [["imbalance", "재료 불균형(퀸 9·룩 5·비숍/나이트 3)", "−0.0442"],
         ["capture_count", "현재 측의 합법 캡처 수", "−0.0152"],
         ["legal_moves", "합법 수 총 개수(분기 인자)", "−0.0202"],
         ["total_material", "양측 비폰 재료 가치 합(게임 단계)", "−0.0229"],
         ["forcing_total", "양측 캡처+체크 수 합", "−0.0457"],
         ["is_check", "현재 체크 상태(0/1)", "−0.0289"],
         ["is_endgame", "재료 ≤ 24(종반 여부)", "+0.0395"]])
    caption(doc, "표 3. 정적 FEN feature와 best_t의 Spearman 순위 상관")
    h2(doc, "5.4 다변량 예측 결과")
    body(doc,
         "여러 feature의 조합·비선형 관계를 포착하기 위해 7종을 모두 입력한 RandomForest 회귀의 "
         "결과는 표 4와 같다. 보지 않은 test에서의 MAE는 10회 평균 1.811로, 동일 예산 균일 정책 "
         "wu(250)의 1.810과 사실상 동일하였다(회귀의 평균 탐색 시간은 243ms). 즉 비선형·상호작용까지 "
         "포착하는 강한 모델을 7종 feature에 모두 투입하여도 균일 정책을 이기지 못하였다.")
    make_table(doc,
        ["정책", "test MAE (10회 평균)"],
        [["RandomForest 회귀(7 feature)", "1.811"],
         ["wu (균일 정책, 250ms)", "1.810"]])
    caption(doc, "표 4. RandomForest 회귀와 균일 정책의 비교 (hold-out test, 10회 평균)")
    body(doc,
         "이 결과의 해석은 다음과 같다. 정적 feature에 예측 신호가 없으므로, 오차를 최소화하려는 "
         "회귀 모델은 모든 국면에 대해 best_t의 평균값에 가까운 값을 예측하게 된다. 그 결과 예측 "
         "탐색 시간이 평균(약 243ms)으로 수렴하여, 정책이 사실상 균일 정책 wu와 동일해진다. 또한 "
         "학습셋에서의 성능이 보지 않은 test보다 소폭 높을 뿐 그 우위가 새로운 국면에서는 사라져, "
         "일반화되는 패턴이 존재하지 않음을 확인하였다.")
    h2(doc, "5.5 종합 해석")
    body(doc,
         "단변량(Spearman)과 다변량(RandomForest) 분석은 일관되게 동일한 결론을 가리킨다: 이론적 "
         "상한이 보여준 38%의 잠재력은 분명히 존재하지만, 정적 FEN 정보만으로는 이에 도달할 수 없다. "
         "특히 인간 기준으로 단순해 보이는 일부 국면에서도 엔진이 체크메이트 수순 탐지에 실패하는 "
         "사례가 관찰되었는데, 이는 ‘사람의 직관 기준 단순한 국면이 반드시 엔진의 요구 연산량이 "
         "낮음을 의미하지는 않는다’는 점을 단적으로 보여준다. 인간의 복잡도 인식과 엔진의 실제 연산 "
         "요구량 사이에는 정적 feature로 메우기 어려운 간극이 존재한다.")

    # ---------------- 6. 결론 ----------------
    h1(doc, "6. 결론")
    body(doc,
         "본 연구는 저사양 클라우드 인스턴스(t3.micro)에서 국면 복잡도에 따라 엔진 탐색 시간을 차등 "
         "배분하여 실시간 형세 평가치의 품질을 높일 수 있는지를 검증하였다. 결과는 다음 두 가지로 "
         "요약된다. 첫째, 각 국면이 자신의 최적 탐색 시간을 부여받는 이론적 상한은 동일 예산 균일 "
         "정책 대비 약 38%의 개선 잠재력을 보여, 차등 배분 전략 자체는 원리적으로 유효하다. 둘째, "
         "그러나 탐색 이전의 정적 FEN 정보만으로는 — 단변량 상관(|ρ| < 0.05)이든 다변량 "
         "RandomForest(균일 정책과 동일한 MAE)든 — 이 잠재력에 도달할 수 없다. 인간 기준의 단순한 "
         "국면이 엔진의 낮은 요구 연산량을 의미하지 않으며, FEN만으로 필요한 연산량을 사전 예측하는 "
         "데에는 분명한 한계가 있다.")
    body(doc,
         "본 연구의 결과는 negative result에 해당하나, 동일 예산 비교와 이론적 상한 실험을 통해 "
         "가설이 성립하지 않는 원인을 ‘차등 배분 전략의 한계’가 아니라 ‘정적 정보의 예측력 한계’로 "
         "분리하여 실험적으로 규명한 데 의의가 있다. 이는 후속 연구가 동일한 방향(정적 FEN 정보 "
         "기반 사전 예측 + 시간 차등 배분)으로 자원을 투입하지 않도록 하는 가이드를 제공한다.")
    body(doc,
         "향후 과제는 다음과 같다. 첫째, 본 연구의 MAE는 모든 국면을 동등 가중으로 평가하지만, "
         "실시간 형세 분석의 사용자 체감 측면에서는 승률 50% 근방의 팽팽한 국면에서의 오차가 더 큰 "
         "영향을 미친다. 이에 가중 MAE 지표를 통해 결론이 사용자 체감 기준에서도 유지되는지 추가로 "
         "검증할 수 있다. 둘째, 본 연구가 확인한 한계는 ‘정적’ 정보에 국한된다. 짧은 예비 탐색에서 "
         "얻는 동적 신호(예: 얕은 깊이에서의 평가치 변동성) 등 탐색 중 정보를 활용하면 best_t 예측 "
         "가능성이 달라질 수 있으므로, 이를 후속 연구 방향으로 제안한다.")

    # ---------------- 참고문헌 ----------------
    h1(doc, "참고문헌")
    refs = [
        '[1] R. B. Miller, "Response time in man-computer conversational transactions," in Proc. AFIPS Fall Joint Comput. Conf., 1968, pp. 267-277.',
        '[2] M. Acher and F. Esnault, "Large-scale Analysis of Chess Games with Chess Engines: A Preliminary Report," arXiv preprint arXiv:1607.04186, 2016.',
        '[3] Chess.com, "Chess.com Reaches 250 Million Members." [Online]. Available: https://www.chess.com/news/view/chesscom-reaches-250-million-members. (Accessed: Mar. 18, 2026)',
        '[4] Chess.com, "An Update Regarding Our Server." [Online]. Available: https://www.chess.com/blog/CHESScom/an-update-regarding-our-server. (Accessed: Mar. 18, 2026)',
        '[5] Chess.com, "Chess.com Gameplay Now Faster With Servers Across World." [Online]. Available: https://www.chess.com/news/view/chesscom-gameplay-now-faster-with-servers-across-world. (Accessed: Mar. 18, 2026)',
        '[6] Lichess.org, "Stockfish.wasm," GitHub. [Online]. Available: https://github.com/lichess-org/stockfish.wasm. (Accessed: Mar. 18, 2026)',
        '[7] Lichess.org, "Evaluation to Win Rate Conversion," lichess-org/scalachess, GitHub. [Online]. Available: https://github.com/lichess-org/scalachess/blob/master/core/src/main/scala/eval.scala#L66',
        '[8] Chess Programming Wiki, "Lazy SMP." [Online]. Available: https://www.chessprogramming.org/Lazy_SMP',
        '[9] Y. Jiang and L. Zhang, "Beyond Accuracy: A Geometric Stability Analysis of LLMs in Chess Evaluation," arXiv:2512.15033, Dec. 2025.',
        '[10] J. Cohen, Statistical Power Analysis for the Behavioral Sciences, 2nd ed. Hillsdale, NJ: Lawrence Erlbaum Associates, 1988.',
        '[11] Computer Chess Rating Lists (CCRL). [Online]. Available: https://computerchess.org.uk/4040/ (Accessed: Mar. 18, 2026)',
    ]
    for r in refs:
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        run(p, r, size=9.5)

    Path(OUT).parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"저장 완료: {OUT}")
    chk = Document(OUT)
    print(f"  문단 {len(chk.paragraphs)}개, 표 {len(chk.tables)}개")


if __name__ == "__main__":
    main()
