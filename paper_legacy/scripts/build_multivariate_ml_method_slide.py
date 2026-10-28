#!/usr/bin/env python3
"""다변량 정적 ML — 방법 보충 슬라이드 (결과 슬라이드 앞에 배치).

RandomForest 회귀로 best_t를 예측하는 파이프라인 + RF 직관 + 용어/검증 설명.
결과 슬라이드(multivariate_ml_slide.pptx)는 건드리지 않고 별도 파일로 생성.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).resolve().parent.parent / "multivariate_ml_method_slide.pptx"

ACCENT = "1F4E79"   # 남색 (표 헤더와 동일 톤)
LIGHT = "EEF2F8"    # 연한 박스
GRAY = "666666"


def add_title(slide, text, size=23):
    title = slide.shapes.title
    title.text = text
    for p in title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string(ACCENT)


def pipeline_box(slide, left, top, w, h, lines, fill=ACCENT, font_color="FFFFFF",
                 font_size=12, bold=True):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(left), Inches(top), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string(fill)
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(font_color)
    return box


def arrow(slide, left, top, w=0.32, h=1.0):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "▶"
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(ACCENT)


def explain_box(slide, left, top, w, h, head, items):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(left), Inches(top), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string(LIGHT)
    box.line.color.rgb = RGBColor.from_string(ACCENT)
    box.line.width = Pt(1.0)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.14)
    p0 = tf.paragraphs[0]
    p0.text = head
    for r in p0.runs:
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(ACCENT)
    for it in items:
        p = tf.add_paragraph()
        p.text = "•  " + it
        p.space_before = Pt(5)
        for r in p.runs:
            r.font.size = Pt(12)
    return box


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    add_title(slide, "방법 — 정적 feature를 RandomForest 회귀에 넣어 적정 탐색시간(best_t) 예측")

    add_textbox = slide.shapes.add_textbox
    sub = add_textbox(Inches(0.5), Inches(0.95), Inches(12.4), Inches(0.5))
    sub.text_frame.word_wrap = True
    sp = sub.text_frame.paragraphs[0]
    sp.text = ("질문: 국면을 보는 것만으로(=정적 정보) \"이 국면엔 몇 ms를 쓰면 좋은가\"를 예측해, "
               "그 시간을 배분하면 평가 품질(MAE)이 균일 배분보다 좋아지는가?")
    for r in sp.runs:
        r.font.size = Pt(12)
        r.font.italic = True

    # ── 파이프라인 (5 박스 + 화살표) ──
    boxes = [
        ["국면", "(FEN)"],
        ["정적 feature", "7개 추출"],
        ["RandomForest", "회귀 학습"],
        ["best_t 예측", "→ 50ms 격자로"],
        ["그 시간의", "평가오차(MAE)"],
    ]
    x = 0.5
    bw, bh, gap = 2.2, 1.05, 0.34
    top = 1.75
    for i, lines in enumerate(boxes):
        pipeline_box(slide, x, top, bw, bh, lines)
        if i < len(boxes) - 1:
            arrow(slide, x + bw + (gap - 0.32) / 2, top + (bh - 1.0) / 2)
        x += bw + gap

    add_label = add_textbox(Inches(0.5), Inches(2.9), Inches(12.4), Inches(0.4))
    al = add_label.text_frame.paragraphs[0]
    al.text = ("※ best_t = 각 국면에서 평가오차가 가장 작아지는 탐색시간 "
               "(50·100·…·500ms 중 하나). 이걸 정적 feature로 맞힐 수 있는가가 핵심.")
    for r in al.runs:
        r.font.size = Pt(11)
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string(GRAY)

    # ── 좌하: RandomForest란? ──
    explain_box(
        slide, 0.5, 3.55, 6.1, 3.0,
        "RandomForest(랜덤 포레스트)란?",
        [
            "여러 개의 결정트리(decision tree)를 만들어 예측을 평균내는 앙상블 모델",
            "각 트리가 feature 값으로 국면을 갈라가며 best_t를 추정 → "
            "비선형·feature 간 상호작용까지 자동으로 포착",
            "수작업 임계값·버킷 규칙보다 강력 → "
            "\"feature에 신호가 조금이라도 있으면 RF가 잡아낸다\"는 강한 검증",
            "설정: 트리 400개, 잎 최소 5개(과적합 억제). best_t를 연속값으로 직접 "
            "예측하는 '회귀'로 사용",
        ],
    )

    # ── 우하: 용어·검증 ──
    explain_box(
        slide, 6.8, 3.55, 6.05, 3.0,
        "어떻게 검증했나",
        [
            "전체 국면을 train 80% / test 20%로 무작위 분할 → "
            "train으로 학습, 안 본 test에서만 성능 측정",
            "분할을 10번 다르게 반복 → \"운 좋은 한 분할\" 결과가 아님을 확인",
            "비교 대상 wu(균일 정책) = 모든 국면에 같은 시간(250ms)을 주는 정책",
            "평가지표 MAE(평균절대오차): 엔진 평가치가 정답에서 얼마나 벗어났나 "
            "(낮을수록 좋음)",
        ],
    )

    prs.save(OUT)
    print(f"저장: {OUT}")


if __name__ == "__main__":
    main()
